"""Builds fixtures/vault from content.py. Run via scripts/make-fixtures.sh.

Needs python-pptx, python-docx, openpyxl, reportlab, pillow (uv provides them),
plus ffmpeg and macOS `say` for media. The 3D model comes from make_lamp.swift.
"""
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
from datetime import datetime, timedelta, timezone

from PIL import Image, ImageDraw, ImageFilter, ImageFont

import content as C

VAULT = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(__file__), "..", "vault"))
FONT = "/System/Library/Fonts/Helvetica.ttc"
MARKER = "/System/Library/Fonts/Supplemental/Chalkboard.ttc"

INK = (0x1F, 0x23, 0x2A)
GREY = (0x5C, 0x63, 0x6E)
CORAL = (0xE8, 0x6A, 0x5C)


def out(rel):
    p = os.path.join(VAULT, rel)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


def log(rel):
    size = os.path.getsize(os.path.join(VAULT, rel))
    print(f"  {rel:<42} {size / 1024:>8.1f} KB")


# ---------------------------------------------------------------- pptx

def make_deck(spec):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height

    def rgb(t):
        return RGBColor(*t)

    def style_title(shape, size=34):
        shape.left, shape.top, shape.width, shape.height = Inches(0.8), Inches(0.55), W - Inches(1.6), Inches(1.0)
        tf = shape.text_frame
        tf.word_wrap = True
        for p in tf.paragraphs:
            p.alignment = 1  # left
            for r in p.runs:
                r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, rgb(INK), "Helvetica Neue"

    def rule(slide):
        line = slide.shapes.add_shape(1, Inches(0.8), Inches(1.55), Inches(0.9), Emu(38100))
        line.fill.solid()
        line.fill.fore_color.rgb = rgb(CORAL)
        line.line.fill.background()

    def footer(slide, n):
        tb = slide.shapes.add_textbox(W - Inches(2.3), H - Inches(0.6), Inches(1.6), Inches(0.3))
        p = tb.text_frame.paragraphs[0]
        p.text = f"{C.COMPANY}  ·  {n}"
        p.alignment = 3  # right
        p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(10), rgb(GREY)

    # Title slide
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = spec["title"]
    s.placeholders[1].text = spec["subtitle"]
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.8), Inches(2.6), W - Inches(1.6), Inches(1.3)
    for r in t.text_frame.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(48), True, rgb(INK), "Helvetica Neue"
    t.text_frame.paragraphs[0].alignment = 1
    sub = s.placeholders[1]
    sub.left, sub.top, sub.width, sub.height = Inches(0.8), Inches(3.9), W - Inches(1.6), Inches(0.8)
    for r in sub.text_frame.paragraphs[0].runs:
        r.font.size, r.font.color.rgb, r.font.name = Pt(20), rgb(GREY), "Helvetica Neue"
    sub.text_frame.paragraphs[0].alignment = 1
    rule(s)
    s.shapes[-1].top = Inches(2.4)

    for i, (title, body, notes) in enumerate(spec["slides"], start=2):
        if body in ("CHART", "TABLE"):
            s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
        else:
            s = prs.slides.add_slide(prs.slide_layouts[1])  # title and content
        s.shapes.title.text = title
        style_title(s.shapes.title)
        rule(s)
        if body == "CHART":
            data = CategoryChartData()
            data.categories = [c for c, _ in C.CHANNEL_REVENUE]
            data.add_series("Q3 revenue ($k)", [v for _, v in C.CHANNEL_REVENUE])
            gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.8), Inches(1.9), Inches(8.5), Inches(4.8), data)
            ch = gf.chart
            ch.has_legend = False
            plot = ch.plots[0]
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(14)
            plot.data_labels.number_format = '"$"0"k"'
            plot.data_labels.number_format_is_linked = False
            plot.series[0].format.fill.solid()
            plot.series[0].format.fill.fore_color.rgb = rgb(CORAL)
            ch.category_axis.tick_labels.font.size = Pt(14)
            ch.value_axis.visible = False
            ch.value_axis.has_major_gridlines = False
            tb = s.shapes.add_textbox(Inches(9.6), Inches(2.4), Inches(3.0), Inches(3))
            tf = tb.text_frame
            tf.word_wrap = True
            for j, line in enumerate(["$412k total", "DTC 61%", "Retail 27%", "B2B 12%"]):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = line
                p.runs[0].font.size = Pt(26 if j == 0 else 18)
                p.runs[0].font.bold = j == 0
                p.runs[0].font.color.rgb = rgb(INK if j == 0 else GREY)
        elif body == "TABLE":
            rows = C.UNIT_ECON_ROWS
            gt = s.shapes.add_table(len(rows), 3, Inches(0.8), Inches(2.0), Inches(9), Inches(0.5) * len(rows))
            tbl = gt.table
            tbl.columns[0].width = Inches(4.4)
            tbl.columns[1].width = tbl.columns[2].width = Inches(2.3)
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = tbl.cell(r, c)
                    cell.text = val
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb((0xF4, 0xF4, 0xF5) if r == 0 else (0xFF, 0xFF, 0xFF))
                    for run in cell.text_frame.paragraphs[0].runs:
                        run.font.size = Pt(16)
                        run.font.bold = r == 0 or r == len(rows) - 1
                        run.font.color.rgb = rgb(INK)
                    if c:
                        cell.text_frame.paragraphs[0].alignment = 3
        else:
            ph = s.placeholders[1]
            ph.left, ph.top, ph.width, ph.height = Inches(0.8), Inches(1.9), W - Inches(1.6), Inches(4.9)
            tf = ph.text_frame
            for j, line in enumerate(body):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.text = line
                p.space_after = Pt(10)
                p.runs[0].font.size = Pt(22)
                p.runs[0].font.color.rgb = rgb(INK)
                p.runs[0].font.name = "Helvetica Neue"
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        footer(s, i)

    prs.core_properties.title = spec["title"]
    prs.core_properties.author = "Priya Raman"
    prs.save(out(spec["path"]))
    log(spec["path"])


# ---------------------------------------------------------------- docx

def make_docx(spec, options=None):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, RGBColor

    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Helvetica Neue", Pt(11)
    for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13)):
        d.styles[name].font.name = "Helvetica Neue"
        d.styles[name].font.size = Pt(size)
        d.styles[name].font.color.rgb = RGBColor(*INK)
    d.add_heading(spec["title"], level=0)
    for k, v in spec.get("meta", []):
        p = d.add_paragraph()
        p.add_run(f"{k}: ").bold = True
        p.add_run(v)
    subsections = {"Light quality", "Mechanics", "Controls"}
    for heading, body in spec["sections"]:
        if heading == "Requirements":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        d.add_heading(heading, level=2 if heading in subsections else 1)
        if body == "TABLE":
            rows = spec["options"]
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Light List Accent 1" if "Light List Accent 1" in [s.name for s in d.styles] else "Table Grid"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    t.cell(r, c).text = val
                    if r == 0:
                        t.cell(r, c).paragraphs[0].runs[0].bold = True
            continue
        if body is None:
            continue
        bulleted = heading in subsections or heading in ("Non-goals", "Open issues", "Risks")
        for para in body:
            d.add_paragraph(para, style="List Bullet" if bulleted else None)
    if spec.get("signoff"):
        p = d.add_paragraph()
        p.add_run(spec["signoff"]).italic = True
    d.core_properties.title = spec["title"]
    d.core_properties.author = "Tom Okafor"
    d.save(out(spec["path"]))
    log(spec["path"])


# ---------------------------------------------------------------- xlsx

def make_q3_model():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    a = {row[0]: row[1] for row in C.MODEL_ASSUMPTIONS[1:]}
    asp, retail_margin, landed = a["Arc ASP"], a["Retail margin"], 84.10
    wb = Workbook()
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="F4F4F5")

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
            ws.cell(1, i + 1).font = head
            ws.cell(1, i + 1).fill = fill
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Assumptions"
    header(ws, list(C.MODEL_ASSUMPTIONS[0]), [28, 14, 34])
    for row in C.MODEL_ASSUMPTIONS[1:]:
        ws.append(list(row))
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        c = ws.cell(r, 2)
        if name in ("Arc ASP", "Arc Mini ASP", "Monthly opex", "Opening cash, Jul 1"):
            c.number_format = '"$"#,##0'
        elif isinstance(c.value, float):
            c.number_format = "0.0%"
        elif isinstance(c.value, int):
            c.number_format = "#,##0"
        c.alignment = Alignment(horizontal="right")

    rev = wb.create_sheet("Revenue")
    header(rev, ["Month", "Arc DTC units", "Arc retail units", "Arc revenue", "Other products", "Total revenue"], [10, 15, 16, 15, 16, 16])
    totals = []
    for i, m in enumerate(C.MONTHS):
        dtc, ret = C.ARC_DTC[i], C.ARC_RETAIL[i]
        arc = round(dtc * asp + ret * asp * (1 - retail_margin))
        total = arc + C.OTHER_REVENUE[i]
        totals.append((dtc + ret, arc, total))
        rev.append([f"{m} 2026", dtc, ret, arc, C.OTHER_REVENUE[i], total])
    for r in range(2, 14):
        for c in (2, 3):
            rev.cell(r, c).number_format = "#,##0"
        for c in (4, 5, 6):
            rev.cell(r, c).number_format = '"$"#,##0'

    cogs = wb.create_sheet("COGS")
    header(cogs, ["Month", "Arc units", "Arc landed cost", "Other products COGS", "Total COGS", "Gross margin"], [10, 12, 16, 20, 14, 14])
    gm = []
    for i, m in enumerate(C.MONTHS):
        units, arc, total = totals[i]
        arc_cogs = round(units * landed)
        other = round(C.OTHER_REVENUE[i] * 0.48)
        g = (total - arc_cogs - other) / total
        gm.append((total, arc_cogs + other))
        cogs.append([f"{m} 2026", units, arc_cogs, other, arc_cogs + other, round(g, 4)])
    for r in range(2, 14):
        cogs.cell(r, 2).number_format = "#,##0"
        for c in (3, 4, 5):
            cogs.cell(r, c).number_format = '"$"#,##0'
        cogs.cell(r, 6).number_format = "0.0%"

    summ = wb.create_sheet("Summary")
    header(summ, ["Quarter", "Revenue", "COGS", "Gross margin", "Opex", "Net", "Ending cash"], [10, 14, 14, 14, 14, 14, 14])
    cash = a["Opening cash, Jul 1"]
    for q in range(4):
        rv = sum(gm[i][0] for i in range(q * 3, q * 3 + 3))
        cg = sum(gm[i][1] for i in range(q * 3, q * 3 + 3))
        opex = a["Monthly opex"] * 3
        net = rv - cg - opex
        if q >= 2:
            cash += net
        summ.append([f"Q{q + 1} 2026", rv, cg, round((rv - cg) / rv, 4), opex, net, cash if q >= 2 else None])
    for r in range(2, 6):
        for c in (2, 3, 5, 6, 7):
            summ.cell(r, c).number_format = '"$"#,##0'
        summ.cell(r, 4).number_format = "0.0%"
    wb.properties.title = "Q3 model"
    wb.properties.creator = "Jun Park"
    p = "finance/q3-model.xlsx"
    wb.save(out(p))
    log(p)


def make_unit_economics():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    head, fill = Font(bold=True), PatternFill("solid", fgColor="F4F4F5")

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
            ws.cell(1, i + 1).font = head
            ws.cell(1, i + 1).fill = fill
        ws.freeze_panes = "A2"

    bom = wb.active
    bom.title = "BOM"
    header(bom, ["Component", "Supplier", "Qty", "Unit cost", "Extended"], [28, 22, 6, 12, 12])
    total = 0
    for name, sup, q, u in C.BOM:
        bom.append([name, sup, q, u, round(q * u, 2)])
        total += q * u
    bom.append(["BOM total", None, None, None, round(total, 2)])
    bom.append(["Freight and duty", None, None, None, C.FREIGHT_DUTY])
    bom.append(["Landed cost", None, None, None, round(total + C.FREIGHT_DUTY, 2)])
    for r in range(2, bom.max_row + 1):
        for c in (4, 5):
            bom.cell(r, c).number_format = '"$"#,##0.00'
    for r in range(bom.max_row - 2, bom.max_row + 1):
        bom.cell(r, 1).font = head
        bom.cell(r, 5).font = head
    assert bom["D9"].value == 14.20 and bom["A9"].value.startswith("LED module")

    pr = wb.create_sheet("Pricing")
    header(pr, list(C.PRICING_TIERS[0]), [22, 10, 14, 22, 14, 10])
    for name, price, landed, fees in C.PRICING_TIERS[1:]:
        if landed is None:
            pr.append([name, price, None, None, None, None])
            continue
        contrib = round(price - landed - fees, 2)
        pr.append([name, price, landed, fees, contrib, round(contrib / price, 4)])
    for r in range(2, pr.max_row + 1):
        for c in (2, 3, 4, 5):
            pr.cell(r, c).number_format = '"$"#,##0.00'
        pr.cell(r, 6).number_format = "0.0%"

    sc = wb.create_sheet("Scenarios")
    header(sc, ["Units"] + [f"Contribution at ${p}" for p in C.SCENARIO_PRICES], [10, 20, 20, 20])
    fees = {229: 18.20, 249: 19.20, 279: 20.70}
    for u in C.SCENARIO_UNITS:
        sc.append([u] + [round(u * (p - 84.10 - fees[p])) for p in C.SCENARIO_PRICES])
    for r in range(2, sc.max_row + 1):
        sc.cell(r, 1).number_format = "#,##0"
        for c in (2, 3, 4):
            sc.cell(r, c).number_format = '"$"#,##0'
    wb.properties.title = "Arc unit economics"
    wb.properties.creator = "Tom Okafor"
    p = "finance/unit-economics.xlsx"
    wb.save(out(p))
    log(p)


# ---------------------------------------------------------------- pdf

def make_pdf(spec, table_key):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ink = colors.Color(*(c / 255 for c in INK))
    grey = colors.Color(*(c / 255 for c in GREY))
    title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=22, leading=28, textColor=ink, spaceAfter=18)
    h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=ink, spaceBefore=6, spaceAfter=10)
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=11, leading=16, textColor=ink, spaceAfter=10)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.setFillColor(grey)
        canvas.drawString(inch, 0.6 * inch, spec["title"])
        canvas.drawRightString(LETTER[0] - inch, 0.6 * inch, f"Page {doc.page}")
        canvas.restoreState()

    flow = [Paragraph(spec["title"], title)]
    for i, (heading, paras) in enumerate(spec["pages"]):
        if i:
            flow.append(PageBreak())
        flow.append(Paragraph(heading, h))
        if paras == "TABLE":
            rows = spec[table_key]
            t = Table(rows, hAlign="LEFT", colWidths=[2.2 * inch] + [1.45 * inch] * (len(rows[0]) - 1))
            t.setStyle(TableStyle([
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 10.5),
                ("FONT", (0, 1), (-1, -1), "Helvetica", 10.5),
                ("TEXTCOLOR", (0, 0), (-1, -1), ink),
                ("LINEBELOW", (0, 0), (-1, 0), 0.75, ink),
                ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.lightgrey),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            flow += [t, Spacer(1, 14)]
            if table_key == "pricing":
                flow.append(Paragraph("Prices are ex-works Shenzhen and fixed for the first twelve months. "
                                      "Volume tiers apply per purchase order, not cumulatively.", body))
            continue
        for p in paras:
            flow.append(Paragraph(p, body))
    path = out(spec["path"])
    doc = SimpleDocTemplate(path, pagesize=LETTER, title=spec["title"], author=C.COMPANY,
                            leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    log(spec["path"])


# ---------------------------------------------------------------- images

def font(size, path=FONT, index=0):
    return ImageFont.truetype(path, size, index=index)


def draw_lamp(w, h, glow=1.0, scale=1.0, bg=(18, 18, 20)):
    """A two-section arm lamp on a desk, drawn at 2x and downsampled."""
    S = 2
    W, H = w * S, h * S
    img = Image.new("RGB", (W, H), bg)
    d = ImageDraw.Draw(img)
    # Background falloff
    for y in range(H):
        t = y / H
        c = tuple(int(bg[i] + (6 + 10 * t) * (1 if i else 1)) for i in range(3))
        d.line([(0, y), (W, y)], fill=c)
    desk_y = int(H * 0.78)
    d.rectangle([0, desk_y, W, H], fill=(34, 32, 30))
    k = scale * S * h / 720
    bx, by = int(W * 0.40), desk_y
    elbow = (bx + int(150 * k), by - int(330 * k))
    head = (bx + int(390 * k), by - int(250 * k))
    # Light pool on the desk and cone, drawn on an overlay
    if glow > 0:
        ov = Image.new("RGBA", (W, H), (0, 0, 0, 0))
        od = ImageDraw.Draw(ov)
        cx, cy = head[0] + int(40 * k), desk_y + int(22 * k)
        for r in range(36, 0, -1):
            a = int(10 * glow * (36 - r) / 36 * 2.2)
            rx, ry = int(r * 13 * k), int(r * 2.6 * k)
            od.ellipse([cx - rx, cy - ry, cx + rx, cy + ry], fill=(255, 196, 120, a))
        cone = [(head[0] - int(40 * k), head[1] + int(40 * k)), (head[0] + int(95 * k), head[1] + int(10 * k)),
                (cx + int(380 * k), desk_y), (cx - int(330 * k), desk_y)]
        od.polygon(cone, fill=(255, 210, 150, int(34 * glow)))
        ov = ov.filter(ImageFilter.GaussianBlur(18 * S))
        img = Image.alpha_composite(img.convert("RGBA"), ov).convert("RGB")
        d = ImageDraw.Draw(img)
    metal, edge = (58, 60, 64), (92, 95, 100)
    # Base
    d.ellipse([bx - int(95 * k), by - int(26 * k), bx + int(95 * k), by + int(14 * k)], fill=metal, outline=edge, width=max(1, int(2 * k)))
    d.rectangle([bx - int(14 * k), by - int(40 * k), bx + int(14 * k), by - int(10 * k)], fill=metal)
    # Arms
    for a, b in (((bx, by - int(30 * k)), elbow), (elbow, head)):
        d.line([a, b], fill=metal, width=int(16 * k))
        d.line([a, b], fill=edge, width=int(4 * k))
    for p in ((bx, by - int(30 * k)), elbow):
        d.ellipse([p[0] - int(14 * k), p[1] - int(14 * k), p[0] + int(14 * k), p[1] + int(14 * k)], fill=edge)
    # Shade: a tilted truncated cone
    hx, hy = head
    shade = [(hx - int(30 * k), hy - int(36 * k)), (hx + int(40 * k), hy - int(30 * k)),
             (hx + int(120 * k), hy + int(46 * k)), (hx - int(60 * k), hy + int(64 * k))]
    d.polygon(shade, fill=(70, 72, 76), outline=edge)
    rim_c = (255, 214, 160) if glow > 0.05 else (50, 50, 52)
    lit = tuple(int(50 + (c - 50) * glow) for c in rim_c)
    d.line([shade[3], shade[2]], fill=lit, width=int(7 * k))
    return img.resize((w, h), Image.LANCZOS)


def make_images():
    p = "design/arc-render.png"
    draw_lamp(1600, 1200, glow=1.0).save(out(p), optimize=True)
    log(p)

    # Moodboard: swatches with material names
    W, H = 1800, 1200
    img = Image.new("RGB", (W, H), (236, 233, 228))
    d = ImageDraw.Draw(img)
    sw = [("Graphite", (52, 54, 58)), ("Bone", (226, 219, 203)), ("Coral", (232, 106, 92)),
          ("Brushed aluminium", (176, 178, 180)), ("Walnut desk", (92, 64, 44)), ("Warm 2200K", (255, 190, 120))]
    pad, gw, gh = 60, (W - 60 * 4) // 3, (H - 60 * 3 - 80) // 2
    for i, (name, col) in enumerate(sw):
        x, y = pad + (i % 3) * (gw + pad), pad + (i // 3) * (gh + pad + 40)
        d.rectangle([x, y, x + gw, y + gh], fill=col)
        if name == "Brushed aluminium":
            for yy in range(y, y + gh, 3):
                d.line([(x, yy), (x + gw, yy)], fill=tuple(min(255, c + (8 if yy % 6 else -6)) for c in col))
        if name == "Walnut desk":
            for j in range(18):
                yy = y + j * gh // 18
                d.arc([x - 200, yy - 40, x + gw + 200, yy + 40], 0, 180, fill=(78, 52, 34), width=3)
        d.text((x, y + gh + 14), name, fill=(60, 60, 60), font=font(30))
    d.text((pad, H - 70), "Arc  ·  materials, October 2", fill=(110, 110, 110), font=font(26))
    p = "design/moodboard.jpg"
    img.save(out(p), quality=88)
    log(p)

    # Whiteboard with handwriting, for OCR
    W, H = 1600, 1100
    img = Image.new("RGB", (W, H), (244, 245, 242))
    d = ImageDraw.Draw(img)
    for y in range(H):  # glare
        g = int(6 * math.sin(y / H * math.pi))
        d.line([(0, y), (W, y)], fill=(244 - g, 245 - g, 242 - g))
    blue, red, black = (34, 64, 150), (190, 40, 40), (30, 30, 34)
    mk = lambda s: font(s, MARKER)
    d.text((90, 70), "PRICE LADDER", fill=black, font=mk(84))
    d.line([(90, 175), (760, 168)], fill=black, width=6)
    rows = [("Mini", "$179", blue), ("Arc", "$249", blue), ("Arc Pro", "$329 ?", red)]
    for i, (n, pr, col) in enumerate(rows):
        y = 240 + i * 150
        d.text((130, y), n, fill=col, font=mk(72))
        d.text((560, y), pr, fill=col, font=mk(72))
    d.text((120, 720), "BOM 71.40  ->  58% @ 249", fill=black, font=mk(56))
    d.text((120, 810), "Maya: \"up to 260\"", fill=blue, font=mk(56))
    d.ellipse([1020, 220, 1460, 560], outline=red, width=6)
    d.text((1080, 330), "anchor on\n$279 ?", fill=red, font=mk(60))
    d.line([(960, 330), (1020, 380)], fill=red, width=6)
    img = img.rotate(-1.2, resample=Image.BICUBIC, fillcolor=(200, 200, 198))
    p = "design/whiteboard-pricing.jpg"
    img.save(out(p), quality=85)
    log(p)


# ---------------------------------------------------------------- media

def duration(path):
    r = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", path],
                       capture_output=True, text=True, check=True)
    return float(r.stdout.strip())


def make_teaser():
    W, H, fps, secs = 1280, 720, 24, 8
    tmp = tempfile.mkdtemp()
    big, small = font(120, FONT, 1), font(44)
    for f in range(fps * secs):
        t = f / fps
        glow = 0 if t < 1 else min(1, (t - 1) / 2) ** 1.6
        img = draw_lamp(W, H, glow=glow, scale=0.9, bg=(10, 10, 12))
        ov = Image.new("RGBA", (W, H), (0, 0, 0, 0))
        od = ImageDraw.Draw(ov)
        a1 = int(255 * max(0, min(1, (t - 3.5) / 1.0)))
        a2 = int(255 * max(0, min(1, (t - 5.0) / 0.8)))
        od.text((W - 470, 200), "Arc", fill=(240, 236, 228, a1), font=big)
        od.text((W - 466, 345), "November", fill=(200, 196, 188, a2), font=small)
        img = Image.alpha_composite(img.convert("RGBA"), ov).convert("RGB")
        img.save(os.path.join(tmp, f"f{f:04d}.png"))
    p = out("media/arc-teaser.mp4")
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-framerate", str(fps), "-i", os.path.join(tmp, "f%04d.png"),
                    "-f", "lavfi", "-i", f"sine=frequency=110:duration={secs}",
                    "-f", "lavfi", "-i", f"sine=frequency=165:duration={secs}",
                    "-filter_complex", f"[1][2]amix=inputs=2,volume=0.25,afade=t=in:d=2,afade=t=out:st={secs - 1.5}:d=1.5[a]",
                    "-map", "0:v", "-map", "[a]", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-crf", "24",
                    "-c:a", "aac", "-b:a", "96k", "-movflags", "+faststart",
                    "-metadata", "title=Arc teaser", p], check=True)
    shutil.rmtree(tmp)
    log("media/arc-teaser.mp4")


def make_interview():
    voices = {"tom": "Reed (English (US))", "maya": "Samantha"}
    tmp = tempfile.mkdtemp()
    parts, marks, t = [], {}, 0.0
    gap = 0.55
    sil = os.path.join(tmp, "gap.wav")
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "anullsrc=r=22050:cl=mono",
                    "-t", str(gap), sil], check=True)
    for i, (who, line, tag) in enumerate(C.INTERVIEW):
        aiff = os.path.join(tmp, f"{i:02d}.aiff")
        wav = os.path.join(tmp, f"{i:02d}.wav")
        subprocess.run(["say", "-v", voices[who], "-r", "185" if who == "tom" else "175", "-o", aiff, line], check=True)
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", aiff, "-ar", "22050", "-ac", "1", wav], check=True)
        if tag:
            marks[tag] = t
        parts += [wav, sil]
        t += duration(wav) + gap
    lst = os.path.join(tmp, "list.txt")
    with open(lst, "w") as f:
        f.writelines(f"file '{p}'\n" for p in parts)
    p = out("research/interview-maya-chen.m4a")
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "concat", "-safe", "0", "-i", lst,
                    "-c:a", "aac", "-b:a", "64k", "-metadata", "title=Interview: Maya Chen",
                    "-metadata", "artist=Kestrel Works", p], check=True)
    shutil.rmtree(tmp)
    log("research/interview-maya-chen.m4a")
    return {k: f"{int(v) // 60}:{int(v) % 60:02d}" for k, v in marks.items()}


# ---------------------------------------------------------------- text files

def make_text():
    for rel, body in C.MARKDOWN_NOTES.items():
        with open(out(rel), "w") as f:
            f.write(body)
        log(rel)
    with open(out("tools/forecast.py"), "w") as f:
        f.write(C.FORECAST_PY)
    log("tools/forecast.py")


def make_state(marks):
    """Sidecars in the CONTRACT format and a trail of two sessions."""
    state = os.path.join(VAULT, ".flamingo")
    notes = os.path.join(state, "notes")
    if os.path.isdir(notes):
        shutil.rmtree(notes)
    commented = []
    for target, fid, comments in C.SIDECARS:
        assert os.path.exists(os.path.join(VAULT, target)), target
        lines = ["---", f"target: {target}", f"id: {fid}", "---"]
        for ts, kind, value, text in sorted(comments, key=lambda c: c[0]):
            value = value.format(**marks) if value else value
            text = text.format(**marks)
            anchor = f" ({kind} {value})" if kind else ""
            lines.append(f"- [{ts}]{anchor} " + text.replace("\n", "\n  "))
            commented.append((ts + ":00", "commented", target, kind, value, text))
        rel = os.path.join(".flamingo", "notes", target + ".md")
        with open(out(rel), "w") as f:
            f.write("\n".join(lines) + "\n")
    print(f"  .flamingo/notes/                           {len(C.SIDECARS)} sidecars")

    tz = timezone(timedelta(hours=C.FIXTURE_UTC_OFFSET_HOURS))
    parse = lambda s: datetime.strptime(s, "%Y-%m-%d %H:%M:%S").replace(tzinfo=tz)
    windows = [(parse(a + ":00"), parse(b + ":00")) for a, b in C.TRAIL_SESSIONS]
    events = list(C.TRAIL) + [e for e in commented if any(a <= parse(e[0]) <= b for a, b in windows)]
    events.sort(key=lambda e: e[0])
    # Same shape TrailStore writes: flat file and anchor strings, UTC timestamps.
    ns = uuid.UUID("6ba7b811-9dad-11d1-80b4-00c04fd430c8")
    first_ids = {}
    with open(out(".flamingo/trail.jsonl"), "w") as f:
        for at, kind, file, ak, av, text in events:
            eid = str(uuid.uuid5(ns, f"{at}|{kind}|{file}|{text}")).upper()
            for i, (a, b) in enumerate(windows):
                if a <= parse(at) <= b:
                    first_ids.setdefault(i, eid)
            ev = {"id": eid, "at": parse(at).astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"), "kind": kind}
            if file:
                ev["file"] = file
            if ak:
                ev["anchor"] = f"{ak} {av}"
            if text:
                ev["text"] = text
            f.write(json.dumps(ev, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(f"  .flamingo/trail.jsonl                      {len(events)} events")
    # Name the first session; leave the second unnamed so both states show.
    names = {first_ids[i]: name for i, name in C.TRAIL_NAMES.items()}
    with open(out(".flamingo/trail-names.json"), "w") as f:
        json.dump(names, f, indent=2, sort_keys=True)
        f.write("\n")


def main():
    print(f"Building {VAULT}")
    os.makedirs(VAULT, exist_ok=True)
    only = set(os.environ.get("ONLY", "").split(",")) - {""}
    want = lambda k: not only or k in only
    if want("office"):
        make_deck(C.BOARD_DECK)
        make_deck(C.LAUNCH_DECK)
        make_docx(C.PRODUCT_BRIEF)
        make_docx(C.PRICING_MEMO)
        make_q3_model()
        make_unit_economics()
        make_pdf(C.SUPPLIER_AGREEMENT, "pricing")
        make_pdf(C.MARKET_SIZING, "method")
    if want("images"):
        make_images()
    if want("media"):
        make_teaser()
    marks = make_interview()  # sidecar time anchors come from the real audio
    make_text()
    make_state(marks)


if __name__ == "__main__":
    main()
