"""Disposable end-to-end acceptance documents. No pre-seeded thoughts or trails."""
from pathlib import Path
from pptx import Presentation
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw, ImageFont

root = Path(__file__).resolve().parents[1] / "build/Daily QA"
project = root / "Research Lab"
project.mkdir(parents=True, exist_ok=True)
(project / "notes.md").write_text("# Daily research\n\nThe experiment budget is 6767.\n\n# Decision\n\nCompare the slide evidence before choosing the next experiment.\n", encoding="utf-8")
(project / "策略.txt").write_text("我们的定价策略在第三季度调整。预算为6767。\n", encoding="utf-8")
(root / "exact.txt").write_text("定价策略", encoding="utf-8")
p = Presentation()
for title, body in [("Daily research", "A source navigation check"), ("Evidence", "ORCHIDBUDGET is 6767. Connect this evidence to the working note."), ("Next experiment", "Use the workbook's Q3 B12 value before deciding.")]:
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
p.save(project / "evidence.pptx")
d = Document()
d.add_heading("Research plan", 1)
d.add_paragraph("Start from the saved evidence.")
d.add_heading("Exact conclusion", 1)
d.add_paragraph("QUOKKACONCLUSION: cross-check the 6767 budget before approval.")
d.save(project / "plan.docx")
w = Workbook()
s = w.active
s.title = "Q3"
s["A1"] = "Experiment budget"
s["B12"] = "ORCHIDCELL6767"
w.save(project / "budget.xlsx")
c = canvas.Canvas(str(project / "scan.pdf"))
c.drawString(72, 720, "Page one: source navigation")
c.showPage()
im = Image.new("RGB", (1500, 800), "white")
draw = ImageDraw.Draw(im)
font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 54)
draw.text((75, 150), "SCANNEDORCHID evidence: budget 6767", fill="black", font=font)
c.drawImage(ImageReader(im), 25, 300, width=550, height=293)
c.showPage()
c.drawString(72, 720, "Page three: the final source")
c.save()
stress = root / "Many Files"
stress.mkdir(exist_ok=True)
for i in range(421):
    (stress / f"file-{i:03}.md").write_text(f"Common indexing fixture {i}.\n", encoding="utf-8")
print(root)
