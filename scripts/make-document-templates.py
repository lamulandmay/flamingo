"""Development-only template generator; shipped app has no Python dependency."""
from pathlib import Path
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches
from docx import Document
from openpyxl import Workbook

out = Path(__file__).resolve().parents[1] / "Sources/FlamingoCore/Templates"
out.mkdir(parents=True, exist_ok=True)
p = Presentation()
p.slide_width = Inches(13.333333)
p.slide_height = Inches(7.5)
p.slides.add_slide(p.slide_layouts[0])
p.core_properties.author = ""
p.core_properties.title = ""
p.core_properties.last_modified_by = ""
p.save(out / "Blank.pptx")
d = Document()
d.add_paragraph()
d.core_properties.author = ""
d.core_properties.last_modified_by = ""
d.save(out / "Blank.docx")
w = Workbook()
w.properties.creator = ""
w.properties.lastModifiedBy = ""
w.save(out / "Blank.xlsx")
